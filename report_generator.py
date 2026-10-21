"""Create the downloadable DOCX comparison report with python-docx.

The report contains exactly the sections shown on screen:
Document A, Document B, Similarities, Differences, Key Details Comparison,
Overall Conclusion, Limitations.
"""

from __future__ import annotations

import io
import re
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from comparison import DISCLAIMER, NOT_FOUND

_XML_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")

STATUS_LABELS = {
    "same": "Same",
    "different": "Different",
    "only_in_a": "Only in A",
    "only_in_b": "Only in B",
}
TYPE_LABELS = {
    "added": "Added", "removed": "Removed", "changed": "Changed", "conflict": "Conflict",
}


def _safe(value: Any) -> str:
    """python-docx raises on control characters, so strip them."""
    return _XML_ILLEGAL.sub("", "" if value is None else str(value))


def _shade(cell, hex_fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    props.append(shd)


def _table(doc, headers: list[str], rows: list[list[str]], widths_cm: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(header)
        run.bold = True
        run.font.size = Pt(10)
        _shade(cell, "D9E2F3")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(_safe(value))
            run.font.size = Pt(9.5)
    from docx.shared import Cm
    for row in table.rows:
        for i, width in enumerate(widths_cm):
            row.cells[i].width = Cm(width)
    doc.add_paragraph()


def _excerpt(text: str, page: str, verified: bool | None) -> str:
    if not text:
        return "-"
    out = f"“{text}”"
    if page:
        out += f" (p. {page})"
    if verified is False:
        out += " [unverified]"
    return out


def build_docx_report(result: dict[str, Any]) -> bytes:
    """Return the report as DOCX bytes (ready for ``st.download_button``)."""
    meta = result["meta"]
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    title = doc.add_heading("AI Document Comparison Report", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info.add_run(
        _safe(f"Generated {meta['generated_at']} · Model: {meta['model']}")
    ).italic = True

    note = doc.add_paragraph()
    run = note.add_run(_safe(DISCLAIMER))
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

    # 1-2. Summaries
    for label, name_key, pages_key, sum_key in (
        ("Document A", "file_name_a", "pages_a", "summary_a"),
        ("Document B", "file_name_b", "pages_b", "summary_b"),
    ):
        doc.add_heading(label, level=1)
        p = doc.add_paragraph()
        p.add_run("File name: ").bold = True
        p.add_run(_safe(f"{meta[name_key]} ({meta[pages_key]} page(s))"))
        doc.add_paragraph(_safe(result[sum_key]))

    # 3. Similarities
    doc.add_heading("Similarities", level=1)
    if result["similarities"]:
        rows = [
            [s["point"],
             _excerpt(s["evidence_a"], s["page_a"], s.get("evidence_a_verified")),
             _excerpt(s["evidence_b"], s["page_b"], s.get("evidence_b_verified"))]
            for s in result["similarities"]
        ]
        _table(doc, ["Similarity", "Evidence - Document A", "Evidence - Document B"], rows, [5.5, 5.5, 5.5])
    else:
        doc.add_paragraph("No similarities were identified from the provided text.")

    # 4. Differences
    doc.add_heading("Differences", level=1)
    if result["differences"]:
        rows = [
            [d["topic"], TYPE_LABELS.get(d["type"], d["type"].title()),
             d["document_a"], d["document_b"],
             "A: " + _excerpt(d["evidence_a"], d["page_a"], d.get("evidence_a_verified"))
             + "\nB: " + _excerpt(d["evidence_b"], d["page_b"], d.get("evidence_b_verified"))]
            for d in result["differences"]
        ]
        _table(doc, ["Topic", "Type", "Document A", "Document B", "Evidence"], rows, [3, 1.8, 3.6, 3.6, 4.5])
    else:
        doc.add_paragraph("No differences were identified from the provided text.")

    # 5. Key details
    doc.add_heading("Key Details Comparison", level=1)
    if result["key_details"]:
        rows = [
            [k["field"], k["document_a"], k["document_b"], STATUS_LABELS.get(k["status"], k["status"])]
            for k in result["key_details"]
        ]
        _table(doc, ["Detail", "Document A", "Document B", "Status"], rows, [3.8, 5.2, 5.2, 2.3])
    else:
        doc.add_paragraph(f"{NOT_FOUND} (no comparable key details were identified).")

    # 6. Conclusion
    doc.add_heading("Overall Conclusion", level=1)
    doc.add_paragraph(_safe(result["conclusion"]))

    # 7. Limitations
    doc.add_heading("Limitations", level=1)
    for item in result["limitations"]:
        doc.add_paragraph(_safe(item), style="List Bullet")

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
