import io

from docx import Document

from report_generator import build_docx_report
from test_comparison import FakeLLM, good_payload, run


def open_report(result):
    return Document(io.BytesIO(build_docx_report(result)))


def headings(doc):
    return [p.text for p in doc.paragraphs if p.style.name.startswith(("Heading", "Title"))]


def test_report_contains_all_required_sections():
    doc = open_report(run(FakeLLM(good_payload())))
    h = headings(doc)
    for section in ("Document A", "Document B", "Similarities", "Differences",
                    "Key Details Comparison", "Overall Conclusion", "Limitations"):
        assert section in h, section


def test_report_matches_displayed_content():
    result = run(FakeLLM(good_payload()))
    doc = open_report(result)
    body = "\n".join(p.text for p in doc.paragraphs)
    table_text = "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert "document_a.pdf" in body and "document_b.pdf" in body
    assert result["summary_a"] in body and result["summary_b"] in body
    assert result["conclusion"] in body
    assert "Annual leave entitlement" in table_text and "24 days" in table_text
    assert "Only in B" in table_text
    assert "not legal advice" in body.lower()


def test_report_handles_empty_sections_and_control_chars():
    payload = good_payload(similarities=[], differences=[], key_details=[],
                           summary_a="Bad char \x00\x0b here")
    doc = open_report(run(FakeLLM(payload)))
    body = "\n".join(p.text for p in doc.paragraphs)
    assert "No similarities" in body and "No differences" in body


def test_unverified_quotes_are_marked_in_report():
    payload = good_payload()
    payload["differences"][0]["evidence_b"] = "invented quote that is not in the pdf"
    doc = open_report(run(FakeLLM(payload)))
    assert "[unverified]" in "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
